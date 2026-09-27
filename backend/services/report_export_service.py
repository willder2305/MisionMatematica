"""Generadores de descargas PDF y XLSX para reportes autorizados."""

from datetime import date
from io import BytesIO
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape, letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


def _texto_seguro(valor):
    """Convierte valores a texto y neutraliza formulas al exportar hojas de cálculo."""
    texto = "" if valor is None else str(valor)
    return f"'{texto}" if texto.startswith(("=", "+", "-", "@")) else texto


def _nombre_archivo(prefijo, extension):
    """Crea un nombre de descarga estable sin incluir rutas internas del servidor."""
    return f"{prefijo}_{date.today().isoformat()}.{extension}"


def generar_xlsx(titulo, columnas, filas, prefijo):
    """Construye una hoja XLSX con cabeceras legibles y datos ya autorizados."""
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Reporte"
    hoja.append([titulo])
    hoja.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(columnas))
    hoja["A1"].font = Font(bold=True, color="FFFFFF", size=14)
    hoja["A1"].fill = PatternFill("solid", fgColor="07599C")
    hoja.append([etiqueta for _, etiqueta in columnas])
    for celda in hoja[2]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="0875C9")
    for fila in filas:
        hoja.append([_texto_seguro(fila.get(clave)) for clave, _ in columnas])
    hoja.freeze_panes = "A3"
    hoja.auto_filter.ref = hoja.dimensions
    for indice, columna in enumerate(hoja.iter_cols(min_row=2), start=1):
        letra = get_column_letter(indice)
        ancho = min(max(max(len(str(celda.value or "")) for celda in columna) + 2, 12), 34)
        hoja.column_dimensions[letra].width = ancho
    salida = BytesIO()
    libro.save(salida)
    salida.seek(0)
    return salida, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", _nombre_archivo(prefijo, "xlsx")


def generar_pdf(titulo, columnas, filas, prefijo):
    """Construye un PDF tabular con la misma información autorizada del reporte."""
    salida = BytesIO()
    documento = SimpleDocTemplate(salida, pagesize=landscape(letter), rightMargin=0.35 * inch, leftMargin=0.35 * inch)
    estilos = getSampleStyleSheet()
    total_columnas = len(columnas)
    ancho = (11 * inch - 0.7 * inch) / max(total_columnas, 1)
    datos = [[Paragraph(escape(_texto_seguro(etiqueta)), estilos["BodyText"]) for _, etiqueta in columnas]]
    for fila in filas:
        datos.append([Paragraph(escape(_texto_seguro(fila.get(clave))), estilos["BodyText"]) for clave, _ in columnas])
    tabla = Table(datos, repeatRows=1, colWidths=[ancho] * total_columnas)
    tabla.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#07599C")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 7),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#B9D4E8")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F3F8FC")]),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    documento.build([Paragraph(escape(titulo), estilos["Title"]), Spacer(1, 0.18 * inch), tabla])
    salida.seek(0)
    return salida, "application/pdf", _nombre_archivo(prefijo, "pdf")


def exportar_reporte(formato, titulo, columnas, filas, prefijo):
    """Devuelve el archivo descargable solicitado o None cuando no hay filas exportables."""
    if not filas:
        return None
    if formato == "xlsx":
        return generar_xlsx(titulo, columnas, filas, prefijo)
    if formato == "pdf":
        return generar_pdf(titulo, columnas, filas, prefijo)
    raise ValueError("Formato de exportación no permitido.")
